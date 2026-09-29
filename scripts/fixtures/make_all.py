# -*- coding: utf-8 -*-
"""Genera los documentos de prueba reales en tests/fixtures/.

- ensayo.docx     Word con título, capítulos, listas, tabla, cita, imagen y salto de página.
- texto.pdf       PDF con texto seleccionable: numeración romana + arábiga, marcadores,
                  encabezados y pies repetidos, guiones de corte, una figura y una página
                  a dos columnas (para el aviso de orden de lectura dudoso).
- escaneado.pdf   PDF solo con imágenes (como un escaneo); la última página está degradada.
- escaneado.json  Texto original de cada página escaneada, para medir el OCR.
- antiguo.doc     Word 97-2003 real (convertido con LibreOffice), para el aviso de conversión.

Requiere: python-docx, reportlab, pillow, pyphen y LibreOffice (solo para el .doc).
"""
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile

from PIL import Image, ImageDraw, ImageFilter, ImageFont

sys.path.insert(0, os.path.dirname(__file__))
from corpus import CHAPTERS, LIST_ITEMS, NUMBERED, QUOTE, SCAN_PAGES, SUBTITLE, TABLE, TITLE  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'tests', 'fixtures')
LIB = '/usr/share/fonts/truetype/liberation'
SERIF = os.path.join(LIB, 'LiberationSerif-Regular.ttf')
SERIF_B = os.path.join(LIB, 'LiberationSerif-Bold.ttf')
SERIF_I = os.path.join(LIB, 'LiberationSerif-Italic.ttf')
LIB_SERIF = SERIF
LIB_SERIF_B = SERIF_B


def illustration(path):
    """Ilustración simple: rollo, códice, página y pantalla."""
    w, h = 1200, 520
    img = Image.new('RGB', (w, h), (246, 240, 228))
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(SERIF, 30)
    labels = ['Rollo', 'Códice', 'Imprenta', 'Pantalla']
    for i, label in enumerate(labels):
        cx = 150 + i * 300
        if i == 0:
            d.rounded_rectangle([cx - 80, 120, cx + 80, 330], radius=18, fill=(214, 190, 150), outline=(120, 90, 60), width=4)
            for k in range(5):
                d.line([cx - 60, 160 + k * 32, cx + 60, 160 + k * 32], fill=(150, 120, 90), width=3)
        elif i == 1:
            d.polygon([(cx - 110, 130), (cx, 150), (cx, 340), (cx - 110, 320)], fill=(250, 246, 236), outline=(90, 70, 50))
            d.polygon([(cx + 110, 130), (cx, 150), (cx, 340), (cx + 110, 320)], fill=(250, 246, 236), outline=(90, 70, 50))
        elif i == 2:
            d.rectangle([cx - 85, 110, cx + 85, 340], fill=(255, 255, 255), outline=(60, 60, 60), width=3)
            for k in range(7):
                d.line([cx - 60, 140 + k * 26, cx + 60, 140 + k * 26], fill=(40, 40, 40), width=4)
        else:
            d.rounded_rectangle([cx - 70, 100, cx + 70, 350], radius=22, fill=(40, 38, 36))
            d.rectangle([cx - 56, 124, cx + 56, 326], fill=(250, 248, 242))
            d.line([cx - 44, 220, cx + 44, 220], fill=(156, 74, 38), width=8)
        tw = d.textlength(label, font=font)
        d.text((cx - tw / 2, 390), label, font=font, fill=(60, 50, 40))
    for i in range(3):
        x = 150 + i * 300 + 120
        d.line([x, 230, x + 55, 230], fill=(156, 74, 38), width=5)
        d.polygon([(x + 55, 220), (x + 70, 230), (x + 55, 240)], fill=(156, 74, 38))
    img.save(path)
    return path


# ——— DOCX ———
def make_docx(img_path):
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    doc.styles['Normal'].font.name = 'Georgia'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_paragraph(TITLE, style='Title')
    doc.add_paragraph(SUBTITLE, style='Subtitle')
    p = doc.add_paragraph('Este documento de prueba combina ')
    p.add_run('negritas').bold = True
    p.add_run(', ')
    p.add_run('cursivas').italic = True
    p.add_run(', listas, una tabla, una cita, una imagen y un salto de página, para verificar que la lectura conserve la estructura del original.')

    for ci, (title, paras) in enumerate(CHAPTERS):
        doc.add_heading(f'Capítulo {ci + 1}. {title}', level=1)
        for pi, text in enumerate(paras):
            doc.add_paragraph(text)
            if ci == 0 and pi == 3:
                doc.add_paragraph(QUOTE, style='Quote')
        if ci == 1:
            doc.add_heading('Recomendaciones tipográficas', level=2)
            for item in LIST_ITEMS:
                doc.add_paragraph(item, style='List Bullet')
            doc.add_paragraph('Para empezar a leer:')
            for item in NUMBERED:
                doc.add_paragraph(item, style='List Number')
            table = doc.add_table(rows=len(TABLE), cols=3)
            table.style = 'Table Grid'
            for r, row in enumerate(TABLE):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
            doc.add_paragraph('Figura 1. Cuatro soportes de la lectura.')
            doc.add_picture(img_path, width=Cm(14))
        if ci == 2:
            doc.add_page_break()
            doc.add_heading('Nota sobre los documentos escaneados', level=2)
            doc.add_paragraph('Esta sección empieza después de un salto de página manual insertado en Word, que la aplicación debe respetar en el modo libro.')
    path = os.path.join(OUT, 'ensayo.docx')
    doc.save(path)
    return path


# ——— PDF con texto ———
def make_text_pdf(img_path):
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A5
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import BaseDocTemplate, Frame, Image as RLImage, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer

    pdfmetrics.registerFont(TTFont('Serif', SERIF))
    pdfmetrics.registerFont(TTFont('Serif-Bold', SERIF_B))
    pdfmetrics.registerFont(TTFont('Serif-Italic', SERIF_I))
    from reportlab.lib.fonts import addMapping

    addMapping('Serif', 0, 0, 'Serif')
    addMapping('Serif', 1, 0, 'Serif-Bold')
    addMapping('Serif', 0, 1, 'Serif-Italic')

    body = ParagraphStyle('body', fontName='Serif', fontSize=9.5, leading=13.5, alignment=TA_JUSTIFY, firstLineIndent=12, hyphenationLang='es_ES', embeddedHyphenation=1, uriWasteReduce=0.3)
    h1 = ParagraphStyle('h1', fontName='Serif-Bold', fontSize=17, leading=21, spaceBefore=6, spaceAfter=14)
    h2 = ParagraphStyle('h2', fontName='Serif-Bold', fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=6)
    title = ParagraphStyle('title', fontName='Serif-Bold', fontSize=22, leading=28, alignment=1, spaceAfter=10)
    sub = ParagraphStyle('sub', fontName='Serif-Italic', fontSize=12, leading=16, alignment=1)
    caption = ParagraphStyle('cap', fontName='Serif-Italic', fontSize=8.5, leading=11, alignment=1, spaceBefore=4, spaceAfter=10)

    W, H = A5
    margin = 16 * mm

    class Doc(BaseDocTemplate):
        chapter = ''

        def afterFlowable(self, flowable):
            if isinstance(flowable, Paragraph) and flowable.style.name in ('h1', 'h2'):
                text = flowable.getPlainText()
                key = f'k{self.page}-{abs(hash(text))}'
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(text, key, level=0 if flowable.style.name == 'h1' else 1, closed=False)
                if flowable.style.name == 'h1':
                    Doc.chapter = text

    def front(canv, doc):
        if doc.page == 1:
            canv.addPageLabel(0, 'ROMAN_LOWER', 1)
            canv.addPageLabel(2, 'ARABIC', 1)

    def running(canv, doc):
        n = doc.page - 2
        canv.saveState()
        canv.setFont('Serif-Italic', 7.5)
        head = TITLE if n % 2 == 0 else (Doc.chapter or TITLE)
        canv.drawCentredString(W / 2, H - 10 * mm, head)
        canv.setFont('Serif', 8)
        canv.drawCentredString(W / 2, 9 * mm, str(n))
        canv.restoreState()

    frame = Frame(margin, margin + 4 * mm, W - 2 * margin, H - 2 * margin - 8 * mm, id='f')
    colw = (W - 2 * margin - 6 * mm) / 2
    left = Frame(margin, margin + 4 * mm, colw, H - 2 * margin - 8 * mm, id='l')
    right = Frame(margin + colw + 6 * mm, margin + 4 * mm, colw, H - 2 * margin - 8 * mm, id='r')
    path = os.path.join(OUT, 'texto.pdf')
    doc = Doc(path, pagesize=A5, title=TITLE, author='Renglón (prueba)')
    doc.addPageTemplates([
        PageTemplate('front', [frame], onPage=front),
        PageTemplate('body', [frame], onPage=running),
        PageTemplate('cols', [left, right], onPage=running),
    ])
    story = [Spacer(1, 50 * mm), Paragraph(TITLE, title), Paragraph(SUBTITLE, sub), PageBreak()]
    story += [Paragraph('Índice', ParagraphStyle('idx', parent=h2))]
    for i, (t, _) in enumerate(CHAPTERS):
        story.append(Paragraph(f'Capítulo {i + 1}. {t}', ParagraphStyle('toc', fontName='Serif', fontSize=10, leading=15)))
    story += [NextPageTemplate('body'), PageBreak()]
    for ci, (t, paras) in enumerate(CHAPTERS):
        if ci == 3:
            story += [NextPageTemplate('cols'), PageBreak()]
        elif ci:
            story += [NextPageTemplate('body'), PageBreak()]
        story.append(Paragraph(f'Capítulo {ci + 1}. {t}', h1))
        for pi, text in enumerate(paras):
            story.append(Paragraph(text, body))
            if ci == 1 and pi == 1:
                story.append(Spacer(1, 4))
                story.append(RLImage(img_path, width=W - 2 * margin, height=(W - 2 * margin) * 520 / 1200))
                story.append(Paragraph('Figura 1. Cuatro soportes de la lectura.', caption))
    doc.build(story)
    return path


# ——— PDF escaneado ———
def wrap(draw, text, font, width):
    words, lines, cur = text.split(), [], ''
    for w in words:
        test = (cur + ' ' + w).strip()
        if draw.textlength(test, font=font) <= width:
            cur = test
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def scanned_page(paras, degrade=False, seed=0):
    rnd = random.Random(seed)
    dpi = 200
    w, h = int(8.27 * dpi), int(11.69 * dpi)
    img = Image.new('L', (w, h), 250)
    d = ImageDraw.Draw(img)
    body = ImageFont.truetype(LIB_SERIF, int(dpi * 12 / 72))
    head = ImageFont.truetype(LIB_SERIF_B, int(dpi * 18 / 72))
    x0, y = int(dpi * 1.0), int(dpi * 1.1)
    width = w - 2 * x0
    for i, para in enumerate(paras):
        font = head if i == 0 else body
        lead = int(font.size * 1.45)
        for line in wrap(d, para, font, width):
            d.text((x0, y), line, font=font, fill=22)
            y += lead
        y += int(lead * 0.6)
    d.text((w / 2 - 20, h - dpi * 0.8), str(seed + 1), font=body, fill=40)
    # Imperfecciones de escáner: leve giro, ruido y compresión.
    img = img.rotate(rnd.uniform(-0.6, 0.6), resample=Image.BICUBIC, fillcolor=250)
    px = img.load()
    for _ in range(w * h // 400):
        px[rnd.randrange(w), rnd.randrange(h)] = rnd.randrange(120, 255)
    if degrade:
        small = img.resize((w // 4, h // 4), Image.BILINEAR).filter(ImageFilter.GaussianBlur(0.9))
        img = small.resize((w, h), Image.BILINEAR)
        img = img.point(lambda v: 150 + v * 0.38)  # poco contraste, tinta gris
    return img.convert('RGB')


def make_scanned_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = os.path.join(OUT, 'escaneado.pdf')
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle('Documento escaneado de prueba')
    tmp = tempfile.mkdtemp()
    for i, paras in enumerate(SCAN_PAGES):
        img = scanned_page(paras, degrade=(i == len(SCAN_PAGES) - 1), seed=i)
        p = os.path.join(tmp, f'p{i}.jpg')
        img.save(p, quality=72)
        c.drawImage(p, 0, 0, width=A4[0], height=A4[1])
        c.showPage()
    c.save()
    shutil.rmtree(tmp)
    with open(os.path.join(OUT, 'escaneado.json'), 'w', encoding='utf-8') as f:
        json.dump({'pages': [' '.join(p) for p in SCAN_PAGES], 'degraded': [len(SCAN_PAGES) - 1]}, f, ensure_ascii=False, indent=2)
    return path


def make_doc(docx_path):
    tmp = tempfile.mkdtemp()
    profile = 'file://' + os.path.join(tmp, 'perfil')
    subprocess.run(['soffice', f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'doc:MS Word 97', '--outdir', tmp, docx_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    out = os.path.join(OUT, 'antiguo.doc')
    shutil.move(os.path.join(tmp, 'ensayo.doc'), out)
    shutil.rmtree(tmp)
    return out


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    tmpimg = os.path.join(tempfile.mkdtemp(), 'figura.png')
    illustration(tmpimg)
    docx = make_docx(tmpimg)
    print('ok', docx)
    print('ok', make_text_pdf(tmpimg))
    print('ok', make_scanned_pdf())
    if shutil.which('soffice'):
        print('ok', make_doc(docx))
    else:
        print('LibreOffice no está instalado: se omite antiguo.doc')
